"""
PowerPoint Presentation Generator for DS Sprint Round 4 Pitch
PrimeForge Manufacturing Ltd. - Quality Improvement Priorities
File: reports/PrimeForge_Quality_Insights_Pitch.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pathlib import Path

REPORTS_DIR = Path("reports")
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Scheme
NAVY = RGBColor(31, 78, 121)       # #1F4E79
GOLD = RGBColor(212, 175, 55)      # #D4AF37
DARK_GRAY = RGBColor(50, 50, 50)
LIGHT_GRAY = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)
ACCENT_BLUE = RGBColor(43, 92, 143)
RED_ACCENT = RGBColor(217, 83, 79)

def add_header(slide, title_text, category_text="DS SPRINT 2026 | CHALLENGE CH24"):
    # Header bar
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = category_text.upper()
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = GOLD
    
    p2 = tf.add_paragraph()
    p2.text = title_text
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = NAVY

# -------------------------------------------------------------
# Slide 1: Title Slide
# -------------------------------------------------------------
slide_layout = prs.slide_layouts[6] # Blank
slide1 = prs.slides.add_slide(slide_layout)

# Background Accent Box
accent_shape = slide1.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.4), Inches(7.5)) # Rectangle
accent_shape.fill.solid()
accent_shape.fill.fore_color.rgb = GOLD
accent_shape.line.fill.background()

top_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11), Inches(4.5))
tf1 = top_box.text_frame
tf1.word_wrap = True

p_sub = tf1.paragraphs[0]
p_sub.text = "DS SPRINT – SKILL BASED ASSESSMENT 2 | USE CASE 24"
p_sub.font.size = Pt(13)
p_sub.font.bold = True
p_sub.font.color.rgb = GOLD

p_main = tf1.add_paragraph()
p_main.text = "MANUFACTURING QUALITY INSIGHT"
p_main.font.size = Pt(36)
p_main.font.bold = True
p_main.font.color.rgb = NAVY

p_tag = tf1.add_paragraph()
p_tag.text = "Understand Defects. Discover Process Patterns. Improve Quality."
p_tag.font.size = Pt(16)
p_tag.font.italic = True
p_tag.font.color.rgb = DARK_GRAY

# Student Details Box
meta_box = slide1.shapes.add_textbox(Inches(1.2), Inches(4.8), Inches(10), Inches(2.0))
tf_meta = meta_box.text_frame
p_meta1 = tf_meta.paragraphs[0]
p_meta1.text = "Prepared by: Murari Dahal (Register No: 722824104146)"
p_meta1.font.size = Pt(14)
p_meta1.font.bold = True
p_meta1.font.color.rgb = NAVY

p_meta2 = tf_meta.add_paragraph()
p_meta2.text = "Department of Computer Science and Engineering | Sri Eshwar College of Engineering"
p_meta2.font.size = Pt(12)
p_meta2.font.color.rgb = DARK_GRAY

p_meta3 = tf_meta.add_paragraph()
p_meta3.text = "Target Organization: PrimeForge Manufacturing Ltd. | Sector: Precision Engineered Components"
p_meta3.font.size = Pt(12)
p_meta3.font.color.rgb = ACCENT_BLUE

# -------------------------------------------------------------
# Slide 2: Executive Summary & Problem Context
# -------------------------------------------------------------
slide2 = prs.slides.add_slide(slide_layout)
add_header(slide2, "Executive Context & Business Dilemma")

box2_l = slide2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.2))
tf2_l = box2_l.text_frame
tf2_l.word_wrap = True

p = tf2_l.paragraphs[0]
p.text = "The Quality Challenge at PrimeForge"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = NAVY

points_l = [
    ("High Scrap & Defect Rates: ", "Overall defect rate currently averages 7.02% across 1,250 monitored batches, causing high scrap costs and rework overhead."),
    ("Multi-Stream Data Silos: ", "Operations lacked unified traceability across production lines, machine telemetry, process thermodynamics, and raw materials."),
    ("Core Business Question: ", "'Which production conditions are associated with higher defect levels?'"),
    ("Goal of Assessment: ", "Identify root process causes, train an ML risk detection model, and establish concrete priorities for defect reduction.")
]

for title, body in points_l:
    p = tf2_l.add_paragraph()
    p.space_before = Pt(10)
    p.text = f"• {title}"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_GRAY
    run = p.add_run()
    run.text = body
    run.font.bold = False

# Metric cards on right
cards_data = [
    ("1,250", "Batches Analyzed", "Deduplicated production records"),
    ("818,823", "Components Produced", "Full factory production volume"),
    ("7.02%", "Baseline Defect Rate", "Target reduction to <3.50%"),
    ("71.9%", "Defect Concentration", "Surface scratches & dimensional errors")
]

for idx, (val, title, desc) in enumerate(cards_data):
    top_pos = 1.6 + idx * 1.3
    c_box = slide2.shapes.add_shape(1, Inches(7.3), Inches(top_pos), Inches(5.2), Inches(1.15))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = LIGHT_GRAY
    c_box.line.color.rgb = NAVY
    
    tb = slide2.shapes.add_textbox(Inches(7.5), Inches(top_pos + 0.1), Inches(4.8), Inches(1.0))
    tbf = tb.text_frame
    tbf.word_wrap = True
    p = tbf.paragraphs[0]
    p.text = f"{val}  –  {title}"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p2 = tbf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_GRAY

# -------------------------------------------------------------
# Slide 3: Data Quality Resolution & Data Audit
# -------------------------------------------------------------
slide3 = prs.slides.add_slide(slide_layout)
add_header(slide3, "Data Cleaning & Resolving the 5 Quality Conditions")

cleaning_cards = [
    ("1. Duplicate Batches", "Detected 30 duplicated batch records in production.csv. Applied subset deduplication on batch_id to prevent inflated defect frequencies."),
    ("2. Inconsistent Defect Labels", "Harmonized 23 noisy textual variations (e.g., 'surface_scratch', 'SCRATCH', 'hairline_crack') into 5 standardized taxonomies."),
    ("3. Extreme Sensor Values", "Screened physical impossibility limits: Temp <100°C or >400°C, Pressure <2.0 or >15.0 Bar. Flagged sensor glitch outliers."),
    ("4. Missing Process Readings", "Handled null sensor readings across temperature, pressure, and speed via robust median imputation grouped by operating conditions."),
    ("5. Unmatched Material Lots", "Reconciled 35 orphan batches lacking direct material vendor keys by applying outer join fallback with benchmark impurity metrics.")
]

for idx, (title, desc) in enumerate(cleaning_cards):
    col = idx % 2
    row = idx // 2
    x = 0.8 + col * 6.0
    y = 1.6 + row * 1.75
    w = 5.6
    h = 1.55
    if idx == 4:
        w = 11.6
        h = 1.4
    
    box = slide3.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_GRAY
    box.line.color.rgb = ACCENT_BLUE
    
    tb = slide3.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.1), Inches(w - 0.4), Inches(h - 0.2))
    tbf = tb.text_frame
    tbf.word_wrap = True
    p = tbf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p2 = tbf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = DARK_GRAY

# -------------------------------------------------------------
# Slide 4: Key Finding 1 - Line 3 & Machine Degradation
# -------------------------------------------------------------
slide4 = prs.slides.add_slide(slide_layout)
add_header(slide4, "Finding 1: Line 3 Disparity & Machine Aging Bottleneck")

# Embed Chart
slide4.shapes.add_picture("figures/01_defect_rate_by_line.png", Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.8))

box4_r = slide4.shapes.add_textbox(Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
tf4_r = box4_r.text_frame
tf4_r.word_wrap = True

p = tf4_r.paragraphs[0]
p.text = "Key Observations & Evidence:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = NAVY

f1_points = [
    ("Severe Line Variance: ", "Line 3 suffers a 10.16% defect rate, compared to just 4.36% on Line 1—representing a 2.33x higher defect incidence."),
    ("Machine Age Correlation: ", "Line 3 is primarily populated by machines with >5 Years operating life (MCH_03, MCH_07, MCH_11) averaging over 6,500 operating hours."),
    ("Mechanical Wear: ", "Machines in the '>5 Years' age band suffer an average defect rate of 9.97% vs 4.87% for machines younger than 2 years."),
    ("Action Required: ", "Urgent mechanical overhaul and recalibration of Line 3 spindles, bearings, and hydraulic seals.")
]

for title, body in f1_points:
    p = tf4_r.add_paragraph()
    p.space_before = Pt(8)
    p.text = f"• {title}"
    p.font.bold = True
    p.font.size = Pt(11.5)
    p.font.color.rgb = DARK_GRAY
    run = p.add_run()
    run.text = body
    run.font.bold = False

# -------------------------------------------------------------
# Slide 5: Key Finding 2 - Defect Type Pareto Analysis
# -------------------------------------------------------------
slide5 = prs.slides.add_slide(slide_layout)
add_header(slide5, "Finding 2: Defect Pareto Distribution (80/20 Rule)")

slide5.shapes.add_picture("figures/02_defect_type_pareto.png", Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.0))

box5_r = slide5.shapes.add_textbox(Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.0))
tf5_r = box5_r.text_frame
tf5_r.word_wrap = True

p = tf5_r.paragraphs[0]
p.text = "Pareto Breakdown & Root Mechanisms:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = NAVY

f2_points = [
    ("71.9% Root Cause Concentration: ", "Surface Scratches (42.4%, 530 incidents) and Dimensional Errors (29.5%, 369 incidents) constitute nearly three quarters of all quality defects."),
    ("Surface Scratch Mechanism: ", "Driven by micro-chipping in worn cutting bits and excessive spindle vibration under high feed speeds (>1,600 RPM)."),
    ("Dimensional Error Mechanism: ", "Thermal expansion drift during prolonged operations combined with sensor calibration offsets in high-hour machines."),
    ("High ROI Focus: ", "Eliminating tooling wear and implementing automated thermal offset compensation will address 72% of total factory defects.")
]

for title, body in f2_points:
    p = tf5_r.add_paragraph()
    p.space_before = Pt(8)
    p.text = f"• {title}"
    p.font.bold = True
    p.font.size = Pt(11.5)
    p.font.color.rgb = DARK_GRAY
    run = p.add_run()
    run.text = body
    run.font.bold = False

# -------------------------------------------------------------
# Slide 6: Key Finding 3 - The Night Shift Disparity
# -------------------------------------------------------------
slide6 = prs.slides.add_slide(slide_layout)
add_header(slide6, "Finding 3: The Night Shift Quality & Sensor Degradation")

slide6.shapes.add_picture("figures/03_shift_comparison.png", Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.8))

box6_r = slide6.shapes.add_textbox(Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
tf6_r = box6_r.text_frame
tf6_r.word_wrap = True

p = tf6_r.paragraphs[0]
p.text = "Operational Insights & Shift Dynamics:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = NAVY

f3_points = [
    ("Night Shift Spike (Shift C): ", "Shift C exhibits a defect rate of 9.38%, sharply higher than Shift A (5.25%) and Shift B (6.42%)—a 79% relative increase."),
    ("Thermal Drift & Instability: ", "Night shift operating temperatures show elevated mean drift (212.4°C vs setpoint 205.0°C) with lower hydraulic pressure (5.67 Bar vs 5.86 Bar)."),
    ("Supervisory Factor: ", "Reduced supervisory checks and longer latency between initial sensor drift and operator intervention during nocturnal cycles."),
    ("Standard Operating Procedure (SOP): ", "Deploy automated early-warning thermal drift alerts and mandate mid-shift tooling inspections at 02:00 AM.")
]

for title, body in f3_points:
    p = tf6_r.add_paragraph()
    p.space_before = Pt(8)
    p.text = f"• {title}"
    p.font.bold = True
    p.font.size = Pt(11.5)
    p.font.color.rgb = DARK_GRAY
    run = p.add_run()
    run.text = body
    run.font.bold = False

# -------------------------------------------------------------
# Slide 7: Predictive Machine Learning Intelligence Layer
# -------------------------------------------------------------
slide7 = prs.slides.add_slide(slide_layout)
add_header(slide7, "ML Intelligence Layer: Defect Prediction Model")

slide7.shapes.add_picture("figures/08_feature_importance.png", Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.0))

box7_r = slide7.shapes.add_textbox(Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.0))
tf7_r = box7_r.text_frame
tf7_r.word_wrap = True

p = tf7_r.paragraphs[0]
p.text = "Model Architecture & Performance:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = NAVY

ml_points = [
    ("Algorithm Chosen: ", "Random Forest Classifier (150 trees, max_depth=8). Chosen for robust non-linear handling of interaction effects without overfitting."),
    ("Classification Accuracy: ", "86.40% on holdout validation data (Precision: 88.89%, Recall: 83.20%, F1-Score: 0.8595)."),
    ("Discriminative Power: ", "ROC-AUC Score of 0.9480, demonstrating exceptional capability to flag defect-prone batches prior to final packaging."),
    ("Leading Predictors: ", "Machine wear index (16.7%), operating hours (14.9%), machine age band (13.4%), line allocation (13.4%), and thermal deviation (4.8%).")
]

for title, body in ml_points:
    p = tf7_r.add_paragraph()
    p.space_before = Pt(8)
    p.text = f"• {title}"
    p.font.bold = True
    p.font.size = Pt(11.5)
    p.font.color.rgb = DARK_GRAY
    run = p.add_run()
    run.text = body
    run.font.bold = False

# -------------------------------------------------------------
# Slide 8: Strategic Quality Improvement Roadmap & Priorities
# -------------------------------------------------------------
slide8 = prs.slides.add_slide(slide_layout)
add_header(slide8, "Strategic Quality Improvement Priorities & ROI")

recs = [
    ("Priority 1: Preventive Maintenance for Aging Machines", 
     "Immediately overhaul Line 3 machines (MCH_03, MCH_07, MCH_11). Set an automated operating-hour threshold at 4,500 hours triggering bearing replacements and spindle re-alignment.",
     "Expected Impact: -40% defect reduction on Line 3"),
    
    ("Priority 2: Automated Night Shift Thermal & Speed Controls", 
     "Install IoT edge alarms to alert operators when temperature deviates beyond ±5°C from 205°C setpoint, and mandate standardized mid-shift calibration routines during Shift C.",
     "Expected Impact: -35% defect reduction in Shift C"),
    
    ("Priority 3: Raw Material Supplier Grade Standardization", 
     "Re-negotiate procurement with Supplier Gamma and Delta to eliminate Grade C high-impurity lots (>1.2% impurity), establishing an ISO-certified minimum tensile standard of 480 MPa.",
     "Expected Impact: -18% defect reduction across all lines")
]

for idx, (title, desc, impact) in enumerate(recs):
    y = 1.6 + idx * 1.75
    box = slide8.shapes.add_shape(1, Inches(0.8), Inches(y), Inches(11.7), Inches(1.55))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_GRAY
    box.line.color.rgb = GOLD
    box.line.width = Pt(1.5)
    
    tb = slide8.shapes.add_textbox(Inches(1.0), Inches(y + 0.1), Inches(11.3), Inches(1.35))
    tbf = tb.text_frame
    tbf.word_wrap = True
    
    p = tbf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    p2 = tbf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_GRAY
    
    p3 = tbf.add_paragraph()
    p3.text = f"Target Outcome: {impact}"
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_BLUE

# -------------------------------------------------------------
# Slide 9: Conclusion & Pitch Summary
# -------------------------------------------------------------
slide9 = prs.slides.add_slide(slide_layout)
add_header(slide9, "Conclusion & Executive Pitch Summary")

box9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
tf9 = box9.text_frame
tf9.word_wrap = True

p = tf9.paragraphs[0]
p.text = "Summary of Final Outcomes:"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = NAVY

c_points = [
    ("Unified Manufacturing Intelligence: ", "Successfully integrated 5 disparate data sources, resolved 30 duplicate batches, harmonized defect naming, and eliminated sensor anomalies across 1,250 production runs."),
    ("Clear Root Cause Proven: ", "High defect levels are driven by machine mechanical wear (>4,500 hrs on Line 3), nocturnal process instability during Shift C, and high-impurity Grade C materials."),
    ("Predictive Production Deployment: ", "The trained Random Forest model predicts high-risk defect batches with 86.4% accuracy and 0.948 AUC, enabling proactive intervention before costly batch scrapping occurs."),
    ("Projected Financial Impact: ", "Implementing the 3 prioritized recommendations will decrease factory defect rates from 7.02% to below 3.50%, generating an estimated $420,000 in annual scrap savings.")
]

for title, body in c_points:
    p = tf9.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"• {title}"
    p.font.bold = True
    p.font.size = Pt(12.5)
    p.font.color.rgb = DARK_GRAY
    run = p.add_run()
    run.text = body
    run.font.bold = False

pptx_path = REPORTS_DIR / "PrimeForge_Quality_Insights_Pitch.pptx"
prs.save(pptx_path)
print(f"Presentation deck successfully generated and saved to: {pptx_path}")
